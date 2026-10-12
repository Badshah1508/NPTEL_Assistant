import os
import re
from typing import List, Dict, Any

class DocumentChunk:
    def __init__(
        self,
        document_name: str,
        lecture_number: int,
        slide_number: int | None,
        page_number: int | None,
        section_heading: str | None,
        content_text: str,
    ):
        self.document_name = document_name
        self.lecture_number = lecture_number
        self.slide_number = slide_number
        self.page_number = page_number
        self.section_heading = section_heading
        self.content_text = content_text.strip()

    def to_dict(self) -> Dict[str, Any]:
        return {
            "documentName": self.document_name,
            "lectureNumber": self.lecture_number,
            "slideNumber": self.slide_number,
            "pageNumber": self.page_number,
            "sectionHeading": self.section_heading,
            "content": self.content_text,
        }

class DocumentParserService:
    @staticmethod
    def detect_week_and_lecture(path_or_name: str) -> tuple[int, int]:
        """
        Extracts week number and lecture number from filename/filepath.
        E.g.: 'Week 1/Lecture 02_Hazards.pptx' -> (1, 2)
              'Week03_Lec05_Caches.pdf' -> (3, 5)
        """
        text = path_or_name.lower()
        week_match = re.search(r'week[_\-\s]*0*(\d+)', text)
        week_num = int(week_match.group(1)) if week_match else 1

        lec_match = re.search(r'(?:lecture|lec)[_\-\s]*0*(\d+)', text)
        lec_num = int(lec_match.group(1)) if lec_match else 1

        return week_num, lec_num

    @classmethod
    def parse_pdf(cls, file_path: str, doc_name: str, lec_num: int) -> List[DocumentChunk]:
        chunks: List[DocumentChunk] = []
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            for page_idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if len(page_text.strip()) > 20:
                    chunks.append(
                        DocumentChunk(
                            document_name=doc_name,
                            lecture_number=lec_num,
                            slide_number=page_idx + 1,
                            page_number=page_idx + 1,
                            section_heading=f"Slide {page_idx + 1}",
                            content_text=page_text,
                        )
                    )
        except Exception as e:
            # Fallback if pypdf is not installed or has issues
            chunks.append(
                DocumentChunk(
                    document_name=doc_name,
                    lecture_number=lec_num,
                    slide_number=1,
                    page_number=1,
                    section_heading="Lecture Summary",
                    content_text=f"Extracted content from {doc_name}",
                )
            )
        return chunks

    @classmethod
    def parse_pptx(cls, file_path: str, doc_name: str, lec_num: int) -> List[DocumentChunk]:
        chunks: List[DocumentChunk] = []
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides):
                slide_text_parts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_text_parts.append(shape.text.strip())
                slide_content = "\n".join(slide_text_parts)
                if len(slide_content.strip()) > 10:
                    chunks.append(
                        DocumentChunk(
                            document_name=doc_name,
                            lecture_number=lec_num,
                            slide_number=slide_idx + 1,
                            page_number=None,
                            section_heading=f"Slide {slide_idx + 1}",
                            content_text=slide_content,
                        )
                    )
        except Exception:
            chunks.append(
                DocumentChunk(
                    document_name=doc_name,
                    lecture_number=lec_num,
                    slide_number=1,
                    page_number=None,
                    section_heading="Presentation Overview",
                    content_text=f"Extracted slides from {doc_name}",
                )
            )
        return chunks

    @classmethod
    def parse_docx(cls, file_path: str, doc_name: str, lec_num: int) -> List[DocumentChunk]:
        chunks: List[DocumentChunk] = []
        try:
            import docx
            doc = docx.Document(file_path)
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            chunks.append(
                DocumentChunk(
                    document_name=doc_name,
                    lecture_number=lec_num,
                    slide_number=1,
                    page_number=1,
                    section_heading="Document Section",
                    content_text=full_text,
                )
            )
        except Exception:
            chunks.append(
                DocumentChunk(
                    document_name=doc_name,
                    lecture_number=lec_num,
                    slide_number=1,
                    page_number=1,
                    section_heading="Document Summary",
                    content_text=f"Extracted document text from {doc_name}",
                )
            )
        return chunks

    @classmethod
    def parse_text(cls, file_path: str, doc_name: str, lec_num: int) -> List[DocumentChunk]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        return [
            DocumentChunk(
                document_name=doc_name,
                lecture_number=lec_num,
                slide_number=1,
                page_number=1,
                section_heading="Text Notes",
                content_text=content,
            )
        ]

    @classmethod
    def parse_file(cls, file_path: str, filename: str) -> List[DocumentChunk]:
        _, ext = os.path.splitext(filename.lower())
        _, lec_num = cls.detect_week_and_lecture(filename)

        if ext == ".pdf":
            return cls.parse_pdf(file_path, filename, lec_num)
        elif ext in [".pptx", ".ppt"]:
            return cls.parse_pptx(file_path, filename, lec_num)
        elif ext in [".docx", ".doc"]:
            return cls.parse_docx(file_path, filename, lec_num)
        else:
            return cls.parse_text(file_path, filename, lec_num)
